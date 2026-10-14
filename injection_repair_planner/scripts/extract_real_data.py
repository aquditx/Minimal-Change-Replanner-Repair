"""Extract the real WK41 plant data from the planner's workbook into data/real_wk41.json.

Sources
  * Injection_Weekly_plan_WK_41_01Oct_ver01.xlsx  (sheets InjPlan, Balance, Status)
  * Man_Page.pdf (item master: machine std output, part weight, colour, resin) - transcribed below

Everything that the workbook does NOT contain (cavity counts, machine-mold compatibility,
changeover hours, priorities, due dates for undated orders) is a documented assumption.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "raw" / "Injection_Weekly_plan_WK_41_01Oct_ver01.xlsx"
OUT = ROOT / "data" / "real_wk41.json"

# Transcribed from Man_Page.pdf: product -> list of (machine, std output pcs/hour, colour)
MAN_PAGE = {
    "52.0BC1RD-CL-G": [("IMP01", 18682, "Clear")],
    "51.0BC1RD-CL-S": [("IMP07", 13642, "Clear")],
    "18.0BC1RD-CL-S": [("IMP07", 21600, "Clear")],
    "34.0BC1EC-CL2-S": [("IMP07", 17875, "Clear")],
    "34.0BC1EC-W2-S": [("IMP07", 17875, "White")],
    "12.7BC1EC-CL-S": [("IMP24", 25412, "Clear")],
    "13.5BC1RD-CL-G": [("IMP07", 31418, "Clear"), ("IMP02", 31418, "Clear")],
    "52.0BC1RD-DB-Z": [("IMP01", 18989, "DarkBlue")],
    "9.4TO2RD-BS6-B": [("IMPS01", 95294, "Blue"), ("IMPS02", 95294, "Blue")],
    "26.6TO2RD-CL-B": [("IMP12", 22153, "Clear")],
    "30.5PC2RD-CLS-BH2": [("IMPS06", 43200, "Clear")],
    "36.0BC1RD-CJ-S": [("IMP19", 20945, "Clear")],
    "29-1.15-2TP-NB-A": [("HYELECT1", 61714, "Blue")],
    "29-1.15-2TP-HG-A": [("HYELECT1", 61714, "Green")],
    "29-1.35-2TWOL-AB-A": [("HYCAP1", 89379, "Blue")],
    "29-1.25-2TP-W1-A": [("HYCAP2", 123420, "White")],
    "29-1.25-2TP-CL-A": [("HYCAP2", 123420, "Clear")],
    "52.0BC1RD-CL-B": [("IMP01", 18682, "Clear")],
    "30.5PC2RD-CLR-V": [("IMPS06", 43200, "Clear")],
    "26.6TO2RD-CL-A": [("IMP12", 25984, "Clear")],
    "30.5PC2RD-CLS-A": [("IMPS06", 43200, "Clear")],
    "45.0PC2RD-IGS-B": [("IMPS07", 29965, "Green")],
    "7.45TO2RD-AJBS-B": [("IMPS03", 99692, "Blue")],
    "8.3TO2RD-LBS-A": [("IMPS04", 99694, "Blue")],
    "12.5TO2RD-CLS-B": [("IMPS05", 86400, "Clear")],
    "45.0PC2RD-CLS-B": [("IMPS07", 29965, "Clear")],
}
RESIN_OF = {  # resin group by product (from workbook columns Q / Balance G), default by neck
    "9.4TO2RD-BS6-B": "SOFT", "26.6TO2RD-CL-B": "WTR", "26.6TO2RD-CL-A": "WTR", "36.0BC1RD-CJ-S": "WTR",
    "30.5PC2RD-CLR-V": "RPET", "7.45TO2RD-AJBS-B": "WTR", "8.3TO2RD-LBS-A": "WTR", "12.5TO2RD-CLS-B": "WTR",
}
CAVITY_CHOICES = [4, 8, 12, 16, 24, 32, 48, 64, 72, 96, 128, 144]
ORIGIN_DAY = 2  # t=0 is Fri 02-Oct-2026 00:00
HORIZON = 192.0  # through end of Fri 09-Oct


def day_of(s):
    if not s:
        return None
    m = re.search(r"(\d{2})-Oct", str(s))
    return int(m.group(1)) - ORIGIN_DAY if m else None


def family_of(machine):
    if machine.startswith("IMPS"):
        return "SIPA"
    if machine.startswith("IMP"):
        return "IMP"
    return "CAP"


def neck_of(product):
    for key in ("BC", "TO", "PC"):
        m = re.match(r"^[\d.]+" + key, product)
        if m:
            return key
    return "CAP" if product.startswith("29-") else "OTHER"


def priority_of(customer, note):
    c = (customer or "").lower()
    if (note or "").strip().lower() == "extra" or c == "sfg":
        return 1
    if any(k in c for k in ("almarai", "health water", "middle east", "international refreshments")):
        return 3
    return 2


def pick_cavities(rate):
    target = rate / 850.0
    return min(CAVITY_CHOICES, key=lambda n: abs(n - target))


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["InjPlan"]
    orders, human = [], {}
    seen = {}
    for r in range(4, ws.max_row + 1):
        wrk, po, fg = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        qty = ws.cell(r, 9).value
        if not po or not qty:
            continue  # placeholder rows ("Will update next week")
        po = str(po)
        start, fin = day_of(ws.cell(r, 11).value), day_of(ws.cell(r, 12).value)
        note = ws.cell(r, 13).value or ""
        cust = ws.cell(r, 10).value or ""
        if po in seen:
            # Same order on two rows: the Thursday row (start before t=0) was already produced before
            # the plan origin, so only the later row is remaining work (360k + 216k = 576k total).
            if seen[po]["start_day"] is not None and seen[po]["start_day"] < 0 and (start or 0) >= 0:
                seen[po]["qty"] = float(qty)
                seen[po]["start_day"] = 0
            continue
        o = {"id": po, "product": fg, "qty": float(qty), "customer": cust, "machine": wrk,
             "start_day": start, "finish_day": fin, "note": note,
             "priority": priority_of(cust, note)}
        seen[po] = o
        orders.append(o)
        human.setdefault(wrk, []).append(po)

    bal = wb["Balance"]  # HDPE cap orders (remaining quantity)
    cap_machine = {"29-1.15-2TP-NB-A": "HYELECT1", "29-1.15-2TP-HG-A": "HYELECT1",
                   "29-1.35-2TWOL-AB-A": "HYCAP1", "29-1.25-2TP-W1-A": "HYCAP2",
                   "29-1.25-2TP-CL-A": "HYCAP2"}  # machine from Status sheet / Man_Page
    for r in range(18, 23):
        fg, qty, cust = bal.cell(r, 2).value, bal.cell(r, 4).value, bal.cell(r, 6).value
        if not fg or not qty:
            continue
        po = "CAP-" + fg
        m = cap_machine[fg]
        orders.append({"id": po, "product": fg, "qty": float(qty), "customer": cust or "", "machine": m,
                       "start_day": 0, "finish_day": None, "note": "caps balance", "priority": 2})
        human.setdefault(m, []).append(po)

    machines = sorted({m for v in MAN_PAGE.values() for m, _, _ in v} | {"ENGEL"})
    mach_list = [{"id": m, "family": family_of(m), "speed": 1.0} for m in machines]
    imp_all = [m for m in machines if family_of(m) == "IMP"]
    sipa_all = [m for m in machines if family_of(m) == "SIPA"]
    cap_all = [m for m in machines if family_of(m) == "CAP" and m != "ENGEL"]

    molds = []
    for prod, rows in MAN_PAGE.items():
        neck = neck_of(prod)
        if neck == "BC":
            extra = imp_all
        elif neck in ("TO", "PC"):
            extra = sipa_all + ["IMP12"]
        else:
            extra = cap_all
        n_twins = 2 if prod == "9.4TO2RD-BS6-B" else 1
        rate = rows[0][1]
        cav = pick_cavities(rate)
        resin = RESIN_OF.get(prod, "HDPE" if neck == "CAP" else "CSD")
        for i in range(n_twins):
            home = [m for m, _, _ in rows]
            comp = sorted(set(home) | set(extra))
            mid = f"M:{prod}" + (f"#{i + 1}" if n_twins > 1 else "")
            molds.append({"id": mid, "product": prod, "cavities": cav, "rate": float(rate),
                          "compatible": comp, "color": rows[0][2], "resin": resin, "neck": neck})

    # initial mold per machine = mold of its first planned order (assumption)
    init = {}
    first_prod = {m: next(o["product"] for o in orders if o["id"] == seq[0]) for m, seq in human.items()}
    for m, prod in first_prod.items():
        init[m] = [k["id"] for k in molds if k["product"] == prod][0]

    out = {
        "name": "WK41 real plant (calibrated)", "origin": "2026-10-02T00:00", "horizon_h": HORIZON,
        "machines": mach_list, "molds": molds, "orders": orders, "human_plan": human, "init_mold": init,
        "status_tue_0700": {"IMP07": "Breakdown", "IMP02": "Running 13.5BC1RD-CL-G", "IMP19": "Material runout today"},
        "tue_breakdown_t": 4 * 24 + 7,
        "assumptions": [
            "Std output (R_Round.StdOut) is treated as pieces per hour; consistent with FinishOn dates for most orders.",
            "Cavity counts are synthetic (std output / ~850 pcs per cavity-hour, snapped to common sizes).",
            "Mold-machine compatibility: home machine(s) from Man_Page plus same-family sisters "
            "(BC molds -> IMP lines; TO/PC molds -> SIPA lines + IMP12; cap molds -> cap machines). "
            "Evidence: 13.5BC mold ran on both IMP07 and IMP02.",
            "Due date = end of the FinishOn day (undated orders: end of horizon); where the human plan itself overruns that day the due date is moved to plan completion + 12h so that the old plan starts on time.",
            "Orders listed on two rows with a Thursday row: the Thursday quantity is treated as already produced.",
            "Priorities: SFG/extra=1, Almarai/Nova/Mena/IRC=3, others=2.",
            "Changeover hours come from the synthetic ChangeoverModel (mold + colour + resin).",
        ],
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, indent=1))
    print(f"wrote {OUT}: {len(mach_list)} machines, {len(molds)} molds, {len(orders)} orders")


if __name__ == "__main__":
    main()
